
from openpyxl import load_workbook


def generate_workbook(
    calc_df,
    subject_code,
    roll_col,
    name_col,
    endterm_col,
    sessional_col,
):

    # -----------------------------
    # Load Template
    # -----------------------------

    wb = load_workbook("templates/25-26_2Ytemplate.xlsx")

    # -----------------------------
    # Sessional + End Term Sheet
    # -----------------------------

    ws = wb["Sessional + End Term Assessment"]

    start_row = 8

    for idx, row in calc_df.iterrows():

        excel_row = start_row + idx

        ws[f"A{excel_row}"] = idx + 1
        ws[f"B{excel_row}"] = row[roll_col]
        ws[f"C{excel_row}"] = row[name_col]
        ws[f"D{excel_row}"] = row[endterm_col]
        ws[f"E{excel_row}"] = row[sessional_col]

    # -----------------------------
    # MID Term 1
    # -----------------------------

    ws_mt1 = None

    for sheet in wb.sheetnames:

        if "MID" in sheet and "1" in sheet:

            ws_mt1 = wb[sheet]

            break

    if ws_mt1 is None:

        raise Exception("MID Term 1 sheet not found.")

    start_row = 7

    for idx, row in calc_df.iterrows():

        excel_row = start_row + idx

        ws_mt1[f"D{excel_row}"] = row["CO1(28)"]
        ws_mt1[f"H{excel_row}"] = row["CO2(28)"]
        ws_mt1[f"L{excel_row}"] = row["CO3_M1(14)"]
        ws_mt1[f"R{excel_row}"] = row["MT1"]

    # -----------------------------
    # MID Term 2
    # -----------------------------

    ws_mt2 = None

    for sheet in wb.sheetnames:

        if "MID" in sheet and "2" in sheet:

            ws_mt2 = wb[sheet]

            break

    if ws_mt2 is None:

        raise Exception("MID Term 2 sheet not found.")

    start_row = 7

    for idx, row in calc_df.iterrows():

        excel_row = start_row + idx

        ws_mt2[f"F{excel_row}"] = row["CO3_M2(14)"]
        ws_mt2[f"J{excel_row}"] = row["CO4(28)"]
        ws_mt2[f"N{excel_row}"] = row["CO5(28)"]
        ws_mt2[f"R{excel_row}"] = row["MT2"]

        # -----------------------------
        # Save Workbook
        # -----------------------------

    output_file = f"output/{subject_code}.xlsx"

    wb.save(output_file)

    return output_file