import streamlit as st
import pandas as pd

st.markdown("""
<h1 style='
text-align:center;
color:whiteblack;
background:#1565C0;
padding:15px;
border-radius:10px;
box-shadow:2px 2px 8px rgba(0,0,0,0.3);
'>
🎓 Attainment Generator
</h1>
""", unsafe_allow_html=True)

st.markdown("""
<style>

.stButton > button {
    background-color: #1565C0;
    color: whiteblack;
    font-weight: bold;
    border-radius: 8px;
    border: none;
    padding: 10px 20px;
}

.stButton > button:hover {
    background-color: #0D47A1;
    color: whiteblack;
}

</style>
""", unsafe_allow_html=True)

subject_code = st.text_input("Subject Code")
with open("templates/sample_file_template.xlsx", "rb") as file:
    st.download_button(
        label="📥 Download Sample Marks Template",
        data=file,
        file_name="Student_Marks_Template.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    st.info("""
📌 Instructions

1. Download the sample marks template.
2. Do not change the column names.
3. Enter marks only in the marks columns.
4. Save the file as .xlsx.
5. Upload it below.
""")
uploaded_file = st.file_uploader(
    "Upload Student Marks File",
    type=["xlsx"]
)

if uploaded_file:

    # Read Excel
    df = pd.read_excel(uploaded_file)
    
    # Clean column names
    df.columns = (
        df.columns
        .str.strip()
        .str.replace("\n", " ", regex=False)
        .str.replace("\r", " ", regex=False)
    )

    st.success("File uploaded successfully")
    st.write("Total Students:", len(df))
    st.subheader("Detected Columns")
    st.write(df.columns.tolist())
    st.subheader("Original Data")
    st.dataframe(df.head())
    # Column Detection
    roll_col = df.columns[0]
    name_col = df.columns[1]
    endterm_col = df.columns[2]
    sessional_col = df.columns[3]
    # Calculations
    
    calc_df = df.copy()
    
    student_count = len(calc_df)

    if student_count > 212:
        st.warning(f"More students found than expected: {student_count}")

    elif student_count < 212:
        st.warning(f"Fewer students found than expected: {student_count}")

    else:
        st.success(f"Correct number of students found: {student_count}")
 # Convert marks columns to numeric
    calc_df[sessional_col] = pd.to_numeric(
    calc_df[sessional_col],
    errors="coerce"
    )

    calc_df[endterm_col] = pd.to_numeric(
    calc_df[endterm_col],
    errors="coerce"
    )

    # Remove rows with invalid marks
    calc_df = calc_df.dropna(
    subset=[sessional_col, endterm_col]
    )
    # Convert sessional (30) to 70
    calc_df["Converted70"] = (
        calc_df[sessional_col] * 70 / 30
    ).round().astype(int)

    # Double total
    calc_df["DoubleTotal"] = (
        calc_df["Converted70"] * 2
    )

    # MT1 and MT2 split
    mt1_list = []
    mt2_list = []

    offsets = [0, -1, 1, -2, 2, 0]

    for i, total in enumerate(calc_df["DoubleTotal"]):

        base = total // 2

    # If total is 140, both MT1 and MT2 must be 70
        if total == 140:
            mt1 = 70
            mt2 = 70

        else:

            mt1 = base + offsets[i % len(offsets)]
            mt2 = total - mt1

        # ---------- Business Rule ----------
        # MT1 should not exceed 70
            if mt1 > 70:
                excess = mt1 - 70
                mt1 = 70
                mt2 = mt2 + excess

            # MT2 should not exceed 70
            if mt2 > 70:
                excess = mt2 - 70
                mt2 = 70
                mt1 = mt1 + excess

        # Final safeguard
            mt1 = min(mt1, 70)
            mt2 = min(mt2, 70)

        mt1_list.append(mt1)
        mt2_list.append(mt2)
    calc_df["MT1"] = mt1_list
    calc_df["MT2"] = mt2_list

    # ----------------------------------
    # CO Splitting
    # ----------------------------------

    co1_list = []
    co2_list = []
    co3_m1_list = []

    offsets = [0, -1, 1, -2, 2, 1]

    for i, mt1 in enumerate(calc_df["MT1"]):

        # ---------- Perfect Score ----------
        if mt1 == 70:
            co1 = 28
            co2 = 28
            co3 = 14

        else:

            # ---------- Decide CO3 ----------
            if mt1 >= 65:
                co3 = 13 + (i % 2)        # 13,14
            elif mt1 >= 55:
                co3 = 11 + (i % 3)        # 11,12,13
            elif mt1 >= 45:
                co3 = 9 + (i % 3)         # 9,10,11
            elif mt1 >= 35:
                co3 = 7 + (i % 3)         # 7,8,9
            else:
                co3 = 5 + (i % 3)         # 5,6,7

            remaining = mt1 - co3

            base = remaining // 2

            co1 = base + offsets[i % len(offsets)]
            co2 = remaining - co1

            # ---------- Limits ----------
            co1 = min(max(co1, 0), 28)
            co2 = min(max(co2, 0), 28)

            # ---------- Balance ----------
            if abs(co1 - co2) > 3:

                avg = remaining // 2
                co1 = avg
                co2 = remaining - avg

            co1_list.append(co1)
            co2_list.append(co2)
            co3_m1_list.append(co3)
        calc_df["CO1(28)"] = co1_list
        calc_df["CO2(28)"] = co2_list
        calc_df["CO3_M1(14)"] = co3_m1_list

    #MT2 CO Split 
    co3_m2_list = []
    co4_list = []
    co5_list = []

    offsets = [0, -1, 1, -2, 2, 1]

    for i, mt2 in enumerate(calc_df["MT2"]):

        # ---------- Perfect Score ----------
        if mt2 == 70:
            co4 = 28
            co5 = 28
            co3_m2 = 14

        else:

            # ---------- Decide CO3 ----------
            if mt2 >= 65:
                co3_m2 = 13 + (i % 2)        # 13,14
            elif mt2 >= 55:
                co3_m2 = 11 + (i % 3)        # 11,12,13
            elif mt2 >= 45:
                co3_m2 = 9 + (i % 3)         # 9,10,11
            elif mt2 >= 35:
                co3_m2 = 7 + (i % 3)         # 7,8,9
            else:
                co3_m2 = 5 + (i % 3)         # 5,6,7

            remaining = mt2 - co3_m2

            base = remaining // 2

            co4 = base + offsets[i % len(offsets)]
            co5 = remaining - co4

            # ---------- Limits ----------
            co4 = min(max(co4, 0), 28)
            co5 = min(max(co5, 0), 28)

            # ---------- Balance ----------
            if abs(co4 - co5) > 3:

                avg = remaining // 2
                co4 = avg
                co5 = remaining - avg

            co3_m2_list.append(co3_m2)
            co4_list.append(co4)
            co5_list.append(co5)

        


        calc_df["CO3_M2(14)"] = co3_m2_list
        calc_df["CO4(28)"] = co4_list
        calc_df["CO5(28)"] = co5_list

    st.subheader("Calculated Marks with CO Split")

    st.dataframe(
    calc_df,
    use_container_width=True,
    height=600
    )
    from openpyxl import load_workbook

    if st.button("Load Template"):

        wb = load_workbook(
        "templates/25-26_2Ytemplate.xlsx"
    )

        st.write("Sheet Names:")
        st.write(wb.sheetnames)  
    if st.button("Generate Attainment Sheet"):

        wb = load_workbook("templates/25-26_2Ytemplate.xlsx")

        ws = wb["Sessional + End Term Assessment"]

        start_row = 8

        for idx, row in calc_df.iterrows():

            excel_row = start_row + idx
            ws[f"A{excel_row}"] = idx + 1     # S.No
            ws[f"B{excel_row}"] = row[roll_col]
            ws[f"C{excel_row}"] = row[name_col]
            ws[f"D{excel_row}"] = row[endterm_col]
            ws[f"E{excel_row}"] = row[sessional_col]
        ws_mt1 = None

        for sheet in wb.sheetnames:
            if "MID" in sheet and "1" in sheet:
                ws_mt1 = wb[sheet]
                break

            if ws_mt1 is None:
                ws_mt1 = wb["MID Term 1"]

        start_row = 7

        for idx, row in calc_df.iterrows():

            excel_row = start_row + idx

            ws_mt1[f"D{excel_row}"] = row["CO1(28)"]
            ws_mt1[f"H{excel_row}"] = row["CO2(28)"]
            ws_mt1[f"L{excel_row}"] = row["CO3_M1(14)"]
            ws_mt1[f"R{excel_row}"] = row["MT1"]
        # MID Term 2 Sheet

        ws_mt2 = wb["MID Term 2"]

        start_row = 7

        for idx, row in calc_df.iterrows():

            excel_row = start_row + idx

            ws_mt2[f"F{excel_row}"] = row["CO3_M2(14)"]
            ws_mt2[f"J{excel_row}"] = row["CO4(28)"]
            ws_mt2[f"N{excel_row}"] = row["CO5(28)"]
            ws_mt2[f"R{excel_row}"] = row["MT2"]           

        output_file = f"output/{subject_code}.xlsx"

        wb.save(output_file)

        st.success(f"Workbook saved to {output_file}") 
        with open(output_file, "rb") as file:
            st.download_button(
            label="Download Workbook",
            data=file,
            file_name=f"{subject_code}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ) 
